"""Tests for timesheet_generator.py — validates Excel structure against SPEC.md."""

import datetime
import os
import sys
from pathlib import Path

import openpyxl
import pytest
import yaml

# Import the module under test
sys.path.insert(0, str(Path(__file__).parent.parent))
from timesheet_generator import (
    load_config,
    build_timesheet,
    HEADER_FILL,
    HEADER_FONT,
    WEEKEND_FILL,
    HOLIDAY_FILL,
    COLUMN_WIDTHS,
    CENTER_ALIGN,
    THIN_BORDER,
)


SAMPLE_CONFIG = {
    "employee": {
        "name": "Mohit Banra",
        "id": "FFE520",
        "client": "Partnr",
        "location": "Gurugram",
    },
    "timesheets": [
        {"name": "Test_Aug", "start": "2026-08-01", "end": "2026-08-31"},
        {"name": "Test_Jul", "start": "2026-07-26", "end": "2026-08-26"},
    ],
    "leaves": [
        {"date": "2026-08-06", "type": "Earned Leave"},
        {"date": "2026-08-27", "type": "Medical leave"},
        {"date": "2026-08-15", "type": "Sick Leave"},
    ],
}


def _write_config(cfg: dict, tmpdir: str) -> str:
    path = os.path.join(tmpdir, "config.yaml")
    with open(path, "w") as f:
        yaml.dump(cfg, f)
    return path


# ---------------------------------------------------------------------------
# Config loading tests
# ---------------------------------------------------------------------------

class TestLoadConfig:
    def test_valid_config(self, tmpdir):
        path = _write_config(SAMPLE_CONFIG, tmpdir)
        cfg = load_config(path)
        assert cfg["employee"]["name"] == "Mohit Banra"
        assert len(cfg["timesheets"]) == 2
        assert len(cfg["leaves"]) == 3

    def test_missing_file(self, tmpdir):
        with pytest.raises(FileNotFoundError):
            load_config(os.path.join(tmpdir, "nonexistent.yaml"))

    def test_missing_key(self, tmpdir):
        bad = {"employee": SAMPLE_CONFIG["employee"]}
        path = _write_config(bad, tmpdir)
        with pytest.raises(ValueError, match="Missing config keys"):
            load_config(path)


# ---------------------------------------------------------------------------
# Sheet structure tests
# ---------------------------------------------------------------------------

class TestSheetStructure:
    @pytest.fixture(autouse=True)
    def _setup(self, tmpdir):
        self.cfg = SAMPLE_CONFIG
        self.outdir = tmpdir

    def test_sheet_name_is_month_short_upper(self):
        wb, path = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        assert wb.active.title == "AUG"

    def test_a1_title_format(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        assert ws["A1"].value == "TIMESHEET FOR THE PERIOD 'AUG-2026'"
        assert "A1:F1" in str(ws.merged_cells)

    def test_employee_info_rows(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        assert ws["A2"].value == "EMPLOYEE NAME"
        assert ws["B2"].value == "Mohit Banra"
        assert ws["A3"].value == "EMPLOYEE ID"
        assert ws["B3"].value == "FFE520"
        assert ws["A4"].value == "CLIENT"
        assert ws["B4"].value == "Partnr"
        assert ws["A5"].value == "LOCATION"
        assert ws["B5"].value == "Gurugram"

    def test_header_row(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        for col, expected in enumerate(["WEEKDAY", "DATE", "LOG-IN TIMINGS", "LOG-OUT TIMINGS", "Total Hours", "Remarks"], start=1):
            cell = ws.cell(row=7, column=col)
            assert cell.value == expected
            assert cell.fill.start_color.rgb == HEADER_FILL.start_color.rgb
            assert cell.font.bold is True
            assert cell.alignment == CENTER_ALIGN
            assert cell.border.left.style == THIN_BORDER.left.style

    def test_column_widths(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        for col_letter, expected_width in COLUMN_WIDTHS.items():
            assert ws.column_dimensions[col_letter].width == expected_width

    def test_mon_to_thu_workday(self):
        """Mon 3 Aug 2026: 9:30 AM - 6:30 PM, 9 hours."""
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        # Row 8 = 01-Aug (Sat), row 9 = Sun, row 10 = Mon 3 Aug
        row = 10
        assert ws.cell(row=row, column=1).value == "Monday"
        assert ws.cell(row=row, column=3).value == "9:30:00 AM"
        assert ws.cell(row=row, column=4).value == "6:30:00 PM"
        assert ws.cell(row=row, column=5).value == 9

    def test_friday_workday(self):
        """Fri 7 Aug 2026: 9:30 AM - 5:30 PM, 8 hours."""
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        # 01-Aug is Sat (row 8), so 07-Aug is Fri (row 14)
        row = 14
        assert ws.cell(row=row, column=1).value == "Friday"
        assert ws.cell(row=row, column=4).value == "5:30:00 PM"
        assert ws.cell(row=row, column=5).value == 8

    def test_weekend_fill_and_remark(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        # 01 Aug 2026 = Saturday → row 8
        for col in range(1, 7):
            cell = ws.cell(row=8, column=col)
            assert cell.fill.start_color.rgb == WEEKEND_FILL.start_color.rgb
        assert ws.cell(row=8, column=6).value == "Weekoff"

    def test_leave_fill_and_remark(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        # 06 Aug 2026 = Earned Leave → row 13
        row = 13
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            assert cell.fill.start_color.rgb == HOLIDAY_FILL.start_color.rgb
        assert ws.cell(row=row, column=6).value == "earned leave"  # lowercased per spec

    def test_total_leaves_row_present(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        # Find the "Total Leaves" row
        total_row = None
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == "Total Leaves":
                total_row = r
                break
        assert total_row is not None, "Total Leaves row not found"
        assert ws.cell(row=total_row, column=2).value == 3  # 3 leave entries

    def test_total_leaves_reasons(self):
        wb, _ = build_timesheet(self.cfg, self.cfg["timesheets"][0])
        ws = wb.active
        total_row = None
        for r in range(1, ws.max_row + 1):
            if ws.cell(row=r, column=1).value == "Total Leaves":
                total_row = r
                break
        reasons = ws.cell(row=total_row, column=3).value
        assert "Earned Leave" in reasons
        assert "Medical leave" in reasons
        assert "Sick Leave" in reasons


# ---------------------------------------------------------------------------
# Date range tests
# ---------------------------------------------------------------------------

class TestDateRange:
    def test_jul_aug_range(self):
        """Timesheet 2: 26 Jul → 26 Aug should span both months correctly."""
        wb, _ = build_timesheet(SAMPLE_CONFIG, SAMPLE_CONFIG["timesheets"][1])
        assert wb.active.title == "JUL"  # based on START month
        # Count rows: 26 Jul to 26 Aug = 32 days → 32 data rows + 1 Total Leaves row
        ws = wb.active
        data_rows = sum(
            1 for r in range(8, ws.max_row + 1)
            if ws.cell(row=r, column=2).value is not None
        )
        # 32 data rows + 1 summary row = 33 non-None cells in column B
        assert data_rows == 33


# ---------------------------------------------------------------------------
# End-to-end generation test
# ---------------------------------------------------------------------------

class TestEndToEnd:
    def test_generates_files_and_can_be_read(self, tmpdir):
        cfg_path = _write_config(SAMPLE_CONFIG, tmpdir)
        cfg = load_config(cfg_path)
        paths = []
        for ts in cfg["timesheets"]:
            _, path = build_timesheet(cfg, ts)
            paths.append(path)
        assert len(paths) == 2
        for p in paths:
            assert os.path.exists(p)
            wb = openpyxl.load_workbook(p)
            assert len(wb.sheetnames) == 1